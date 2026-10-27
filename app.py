import os, shutil, subprocess, tempfile, uuid
from pathlib import Path
from flask import Flask, render_template, request, jsonify, send_file
import openpyxl
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

APP_DIR = Path(__file__).resolve().parent
TEMPLATE_XLSX = APP_DIR / "Design of MS-Pipe.xlsx"
OUTPUT_DIR = APP_DIR / "generated"
OUTPUT_DIR.mkdir(exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024

# Input fields are written into the original workbook's Details sheet.
FIELDS = [
    {"key":"project_name","label":"Project name","cell":"B4","type":"text","required":True},
    {"key":"company_name","label":"Company name","cell":"B5","type":"text"},
    {"key":"client_name","label":"Client name","cell":"B6","type":"text"},
    {"key":"mobile","label":"Mobile number","cell":"B7","type":"text"},
    {"key":"email","label":"Email ID","cell":"B8","type":"email"},
    {"key":"discharge","label":"Design discharge (Q)","unit":"m³/s","cell":"B11","type":"number","required":True,"min":0.001,"step":"any"},
    {"key":"length","label":"Total pipe length (L)","unit":"km","cell":"B12","type":"number","required":True,"min":0,"step":"any"},
    {"key":"static_head","label":"Static head (Hs)","unit":"m","cell":"B13","type":"number","required":True,"min":0,"step":"any"},
    {"key":"thickness","label":"Assumed pipe thickness","unit":"mm","cell":"B15","type":"number","required":True,"min":1,"step":"any"},
    {"key":"lining","label":"Internal lining thickness","unit":"mm","cell":"B16","type":"number","min":0,"step":"any"},
    {"key":"guniting","label":"External guniting thickness","unit":"mm","cell":"B17","type":"number","min":0,"step":"any"},
    {"key":"steel_grade","label":"Steel grade","cell":"B19","type":"select","options":["Fe-410","Fe-430","Fe-440","Fe-490","Fe-540","Fe-570","Fe-650","Fe-730","Fe-780"]},
    {"key":"elasticity","label":"Modulus of elasticity of pipe material","unit":"kg/cm²","cell":"B20","type":"number","min":1,"step":"any"},
    {"key":"mortar_density","label":"Density of lining mortar","unit":"kg/cm³","cell":"B21","type":"number","min":0,"step":"any"},
    {"key":"mortar_modulus","label":"Modulus of elasticity of C.M. lining","unit":"kg/cm²","cell":"B22","type":"number","min":0,"step":"any"},
    {"key":"welding","label":"Type of welding","cell":"B24","type":"select","options":["Shop Welding","Field Welding"]},
    {"key":"load_spread","label":"Angle of uniformity spread load","unit":"degrees","cell":"B26","type":"select","options":["0","30","60","90","120","150","180"]},
    {"key":"bedding_angle","label":"Bedding reaction angle","unit":"degrees","cell":"B27","type":"select","options":["0","30","60","90","120","150","180"]},
    {"key":"soil_modulus","label":"Modulus of passive resistance of soil (e′)","unit":"kg/cm²","cell":"B28","type":"number","min":0,"step":"any"},
]

DEFAULTS = {
    "project_name":"MS Pipe Design",
    "company_name":"","client_name":"","mobile":"","email":"",
    "discharge":29.7,"length":100,"static_head":0,"thickness":25,
    "lining":0,"guniting":0,"steel_grade":"Fe-410","elasticity":2100000,
    "mortar_density":0,"mortar_modulus":357000,"welding":"Field Welding",
    "load_spread":"180","bedding_angle":"30","soil_modulus":80
}

def libreoffice_recalculate(input_path, output_dir):
    """Recalculate an XLSX workbook with LibreOffice on Windows/macOS/Linux."""
    import os
    import shutil
    import subprocess
    import tempfile
    from pathlib import Path

    input_path = Path(input_path).resolve()
    output_dir = Path(output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        r"C:\Program Files\LibreOffice\program\soffice.com",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.com",
    ]
    office = next(
        (str(Path(candidate)) for candidate in candidates
         if candidate and Path(candidate).is_file()),
        None,
    )

    if not office:
        raise RuntimeError(
            "LibreOffice was not found. Install LibreOffice, then restart the website. "
            "If it is already installed, update the LibreOffice path in app.py. "
            "Typical location: C:\\\\Program Files\\\\LibreOffice\\\\program\\\\soffice.exe"
        )

    profile_dir = Path(tempfile.mkdtemp(prefix="lo_profile_")).resolve()
    profile_uri = profile_dir.as_uri()
    expected = output_dir / input_path.name

    # LibreOffice may refuse to overwrite an existing output file.
    if expected.exists():
        expected.unlink()

    cmd = [
        office,
        "-env:UserInstallation=" + profile_uri,
        "--headless",
        "--convert-to", "xlsx",
        "--outdir", str(output_dir),
        str(input_path),
    ]

    try:
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
        )
    except FileNotFoundError as exc:
        raise RuntimeError(
            f"Windows could not launch LibreOffice at: {office}. "
            "Check that this file exists and that LibreOffice is installed."
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            "LibreOffice took too long to recalculate the workbook. "
            "Close other LibreOffice windows and try again."
        ) from exc

    if result.returncode != 0 or not expected.exists():
        details = "\n".join(part for part in [result.stdout, result.stderr] if part).strip()
        raise RuntimeError(
            "LibreOffice could not recalculate the workbook. "
            + (details or f"Exit code: {result.returncode}")
        )

    return expected

def clean_value(value):
    if value is None:
        return ""
    if isinstance(value, float):
        return round(value, 4)
    return value

def run_design(form_values):
    job_id = uuid.uuid4().hex[:12]
    job_dir = OUTPUT_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    working = job_dir / "input.xlsx"
    shutil.copy2(TEMPLATE_XLSX, working)

    wb = openpyxl.load_workbook(working)
    ws = wb["Details"]
    for field in FIELDS:
        key, cell = field["key"], field["cell"]
        raw = form_values.get(key, DEFAULTS.get(key, ""))
        if field["type"] == "number":
            if raw in ("", None):
                raw = DEFAULTS.get(key, 0)
            try:
                value = float(raw)
            except (TypeError, ValueError):
                raise ValueError(f"{field['label']} must be a number.")
            if value < field.get("min", 0):
                raise ValueError(f"{field['label']} must be at least {field.get('min', 0)}.")
            ws[cell] = value
        elif key in ("load_spread", "bedding_angle"):
            ws[cell] = int(raw)
        else:
            ws[cell] = raw if raw is not None else ""
    wb.save(working)
    wb.close()

    # LibreOffice recalculates linked formulas from the workbook template.
    recalculated_dir = job_dir / "recalc"
    recalculated_dir.mkdir()
    recalculated = libreoffice_recalculate(working, recalculated_dir)
    final_xlsx = job_dir / "MS_Pipe_Design_Calculated.xlsx"
    shutil.copy2(recalculated, final_xlsx)

    formulas_wb = openpyxl.load_workbook(final_xlsx, data_only=True)
    report_ws = formulas_wb["Final Report"]
    details_ws = formulas_wb["Details"]
    rows = []
    for r in range(1, report_ws.max_row + 1):
        label = report_ws.cell(r, 1).value
        if label is None:
            continue
        unit = report_ws.cell(r, 2).value
        value = report_ws.cell(r, 3).value
        rounded = report_ws.cell(r, 4).value
        # Include headings and rows with a calculated value; omit empty formatting rows.
        if value is not None or (isinstance(label, str) and (label.startswith("(") or label in
            ["Detailed Design Report", "Design of MS-Pipeline", "Diameter Calculation",
             "Total Head-loss Calculation", "Velocity Check"])):
            rows.append({"row":r, "label":str(label), "unit":str(unit or ""),
                         "value":clean_value(value), "rounded":clean_value(rounded)})
    summary = {
        "project_name": details_ws["B4"].value or "MS Pipe Design",
        "company_name": details_ws["B5"].value or "",
        "client_name": details_ws["B6"].value or "",
        "mobile": details_ws["B7"].value or "",
        "email": details_ws["B8"].value or "",
        "discharge": clean_value(report_ws["C13"].value),
        "length": clean_value(report_ws["C14"].value),
        "head": clean_value(report_ws["C28"].value),
        "diameter": clean_value(report_ws["C23"].value),
        "pipe_count": clean_value(report_ws["C22"].value),
        "thickness": clean_value(report_ws["C38"].value),
        "checks": [
            {"name":"Shell thickness","value":report_ws["C81"].value},
            {"name":"Deflection under external load / partial vacuum","value":report_ws["C131"].value},
            {"name":"Deflection under external load with internal pressure","value":report_ws["C138"].value},
            {"name":"Compressive stress","value":report_ws["C148"].value},
            {"name":"Tensile stress","value":report_ws["C156"].value},
            {"name":"Velocity","value":report_ws["B33"].value},
        ],
        "rows": rows,
        "xlsx": str(final_xlsx),
        "job_dir": str(job_dir),
    }
    formulas_wb.close()
    pdf_path = job_dir / "MS_Pipe_Design_Report.pdf"
    create_pdf(summary, pdf_path)
    summary["pdf"] = str(pdf_path)
    return summary

def create_pdf(data, path):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#12304A"), fontSize=19, leading=23, spaceAfter=5))
    styles.add(ParagraphStyle(name="Subtle", parent=styles["Normal"], alignment=TA_CENTER,
                              textColor=colors.HexColor("#526477"), fontSize=9, leading=12))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], textColor=colors.HexColor("#12304A"),
                              fontSize=11, leading=14, spaceBefore=10, spaceAfter=5))
    doc = SimpleDocTemplate(str(path), pagesize=A4, rightMargin=15*mm, leftMargin=15*mm,
                            topMargin=14*mm, bottomMargin=14*mm)
    story = [
        Paragraph("MS PIPELINE DESIGN", styles["ReportTitle"]),
        Paragraph("Detailed design report generated from the supplied Excel calculation workbook", styles["Subtle"]),
        Spacer(1, 7*mm),
    ]
    project_data = [
        ["Project", str(data["project_name"]), "Company", str(data["company_name"] or "—")],
        ["Client", str(data["client_name"] or "—"), "Contact", str(data["mobile"] or "—")],
        ["Email", str(data["email"] or "—"), "Generated", __import__("datetime").datetime.now().strftime("%d %b %Y, %H:%M")],
    ]
    t = Table(project_data, colWidths=[22*mm, 58*mm, 22*mm, 68*mm])
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#F2F6FA")),
        ("TEXTCOLOR",(0,0),(0,-1),colors.HexColor("#526477")),
        ("TEXTCOLOR",(2,0),(2,-1),colors.HexColor("#526477")),
        ("FONTNAME",(0,0),(-1,-1),"Helvetica"),
        ("FONTSIZE",(0,0),(-1,-1),8),
        ("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),
        ("FONTNAME",(2,0),(2,-1),"Helvetica-Bold"),
        ("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D6E0E8")),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LEFTPADDING",(0,0),(-1,-1),6),("RIGHTPADDING",(0,0),(-1,-1),6),
        ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story += [t, Paragraph("Design summary", styles["Section"])]
    summary_data = [
        ["Design discharge", f'{data["discharge"]} m³/s', "Total pipe length", f'{data["length"]} km'],
        ["Total head loss / head", f'{data["head"]} m', "Diameter of each pipe", f'{data["diameter"]} m'],
        ["Number of pipes", str(data["pipe_count"]), "Assumed shell thickness", f'{data["thickness"]} mm'],
    ]
    t2 = Table(summary_data, colWidths=[42*mm, 45*mm, 42*mm, 41*mm])
    t2.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(0,-1),colors.HexColor("#F2F6FA")),
        ("BACKGROUND",(2,0),(2,-1),colors.HexColor("#F2F6FA")),
        ("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"),("FONTNAME",(2,0),(2,-1),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),8),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D6E0E8")),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),6),
        ("TOPPADDING",(0,0),(-1,-1),7),("BOTTOMPADDING",(0,0),(-1,-1),7),
    ]))
    story += [t2, Paragraph("Design checks", styles["Section"])]
    check_data = [["Check", "Workbook result"]]
    for check in data["checks"]:
        check_data.append([str(check["name"]), str(check["value"] if check["value"] not in (None, "") else "Not available")])
    t3 = Table(check_data, colWidths=[115*mm, 55*mm], repeatRows=1)
    t3.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12304A")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),8),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#D6E0E8")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F9FB")]),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),6),
        ("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),
    ]))
    story += [t3, Paragraph("Detailed calculation values", styles["Section"])]
    detail_data = [["Parameter", "Unit", "Calculated value", "Rounded value"]]
    for row in data["rows"]:
        label = row["label"]
        if label in ("Detailed Design Report", "Design of MS-Pipeline"):
            continue
        value = row["value"]
        rounded = row["rounded"]
        if value == "" and rounded == "":
            continue
        detail_data.append([Paragraph(label[:110], styles["Normal"]), row["unit"][:20],
                            str(value)[:35], str(rounded)[:35]])
    t4 = Table(detail_data, colWidths=[83*mm, 24*mm, 32*mm, 31*mm], repeatRows=1)
    t4.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#12304A")),
        ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("FONTSIZE",(0,0),(-1,-1),7),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#D6E0E8")),
        ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F7F9FB")]),
        ("VALIGN",(0,0),(-1,-1),"TOP"),("LEFTPADDING",(0,0),(-1,-1),4),
        ("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),
        ("BOTTOMPADDING",(0,0),(-1,-1),4),
    ]))
    story += [t4, Spacer(1, 5*mm),
              Paragraph("Engineering note: This report reproduces calculations from the provided workbook. Review the input assumptions, applicable design standards, and all results with a qualified engineer before construction or procurement.", styles["Italic"])]
    doc.build(story)

@app.route("/")
def index():
    return render_template("index.html", fields=FIELDS, defaults=DEFAULTS)

@app.post("/api/design")
def api_design():
    values = request.get_json(silent=True) or request.form.to_dict()
    try:
        result = run_design(values)
        # Do not expose server filesystem paths to the browser.
        public = {k:v for k,v in result.items() if k not in ("xlsx","pdf","job_dir","rows")}
        public["downloads"] = {"pdf":f'/download/{Path(result["pdf"]).parent.name}/pdf',
                               "xlsx":f'/download/{Path(result["xlsx"]).parent.name}/xlsx'}
        return jsonify({"ok":True,"result":public})
    except Exception as exc:
        app.logger.exception("Design generation failed")
        return jsonify({"ok":False,"error":str(exc)}), 400

@app.get("/download/<job_id>/<kind>")
def download(job_id, kind):
    job_dir = (OUTPUT_DIR / job_id).resolve()
    if OUTPUT_DIR.resolve() not in job_dir.parents:
        return "Invalid job", 400
    if kind == "pdf":
        file_path = job_dir / "MS_Pipe_Design_Report.pdf"
        mimetype, name = "application/pdf", "MS_Pipe_Design_Report.pdf"
    elif kind == "xlsx":
        file_path = job_dir / "MS_Pipe_Design_Calculated.xlsx"
        mimetype, name = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "MS_Pipe_Design_Calculated.xlsx"
    else:
        return "Invalid file type", 404
    if not file_path.exists():
        return "Report not found. Generate the design again.", 404
    return send_file(file_path, mimetype=mimetype, as_attachment=True, download_name=name)

if __name__ == "__main__":
    # Local development. Production hosting should use the included Gunicorn command.
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False)
